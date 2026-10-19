"""Extract campaign aggregates only. Never exports buyer or contact sheets.
Usage: python scripts/build_data.py workbook.xlsx catalog.json output.json
Requires openpyxl for read-only extraction.
"""
import sys,json,re,unicodedata,collections,datetime,pathlib
import openpyxl
from consolidated import attach_consolidated

def norm(s):
    return re.sub('[^A-Z0-9]','',unicodedata.normalize('NFKD',str(s or '').upper()).encode('ascii','ignore').decode())
def token(s):
    return re.sub(r'[^A-Z0-9]+','_',unicodedata.normalize('NFKD',str(s or '').upper()).encode('ascii','ignore').decode()).strip('_')
def city(s):return 'MDE' if any(x in norm(s) for x in ['MEDELL','MDEL','MDE']) else 'BTA'
def num(v):return float(v) if isinstance(v,(int,float)) and not isinstance(v,bool) else 0
aliases={'BULEVARVERDEAPARTAMENTOS':'BULEVARVERDE','VIVOPARK1':'VIVOPARK','VIVOPARK2':'VIVOPARK','TORREESTACIONCC':'TORREESTACIONCENTROCOMERCIAL','URBANYOFICINAS':'URBANITYOFFICE','URBANYOFI':'URBANITYOFFICE','URBANITYOFICINAS':'URBANITYOFFICE','NUEVATORREESKALA':'ESKALA','TERRA':'URBANIATERRA','TROPICO':'TROPICO'}
def key(n):return aliases.get(norm(n),norm(n))

def build(workbook,catalog_path,output):
    w=openpyxl.load_workbook(workbook,read_only=True,data_only=True)
    catalog=json.loads(pathlib.Path(catalog_path).read_text())
    projects={}; templates=[]; invalid_names=[]
    def add(label,c,**kwargs):
        k=c+'-'+key(label)
        if k not in projects:projects[k]={'id':k,'name':label.title(),'city':c,'token':token(label),'kind':'NO VIS','zone':'','image':None,'url':None,'active':False,'catalog':False,'templates':{}}
        projects[k].update(kwargs);return k
    for c in catalog:
        k=add(c['name'],city(c['city']),name=c['name'],kind=c['kind'],zone=c['zone'],image=c['image'],url=c['url'],logo=c['logo'],catalog=True,available=c['available'])
        if key(c['name'])=='VIVOPARK': projects[k].update(name='Vivopark',token='VIVOPARK')
        if key(c['name'])=='BULEVARVERDE':projects[k].update(name='Bulevar Verde',token='BULEVAR_VERDE')
    for r in list(w['Lista de proyectos X CIUDAD'].values)[3:]:
        for offset,c in [(1,'BTA'),(7,'MDE')]:
            if len(r)<=offset+3 or not isinstance(r[offset],str) or not isinstance(r[offset+3],bool):continue
            k=add(r[offset],c)
            projects[k]['zone']=projects[k]['zone'] or str(r[offset+1] or '')
            projects[k]['kind']='NO VIS' if 'NO VIS' in str(r[offset+2]) else 'VIS'
            projects[k]['premium']='High Ticket' in str(r[offset+2])
    sheets=['Flow Bogota Octubre','Flow Medellín Octubre','Flow Bogota Septiembre',' Flow Medellín Septiembre']
    for sn in sheets:
        c='MDE' if 'Medell' in sn else 'BTA'; rows=list(w[sn].values)
        hi=next(i for i,r in enumerate(rows) if r and r[0]=='Mes')
        for ri,r in enumerate(rows[hi+1:],hi+2):
            if len(r)<16 or not isinstance(r[8],str) or not isinstance(r[10],str) or not r[10].startswith(('FB_','GLG_','TT_','YT_')):continue
            if norm(r[8]) in ['MARCA','SEGUIDORES']:continue
            pid=add(r[8],c); p=projects[pid]
            if 'Octubre' in sn:p['active']=True
            name=r[10].strip(); ch={'FB':'meta','GLG':'google','TT':'tiktok','YT':'youtube'}[name.split('_')[0]]
            if key(r[8]) not in norm(name) and not any(a in norm(name) and b==key(r[8]) for a,b in aliases.items()):
                invalid_names.append({'sheet':sn,'row':ri,'project':p['name'],'name':name,'reason':'El nombre no coincide con el proyecto de la fila.'});continue
            p['kind']='NO VIS' if 'NO_VIS' in name else ('VIS' if '_VIS_' in name else p['kind'])
            item={'name':name,'source':sn,'row':ri}
            if ch not in p['templates']:p['templates'][ch]=item
            if not any(t['project']==pid and t['name']==name for t in templates):templates.append({'project':pid,'channel':ch,**item})
    # Launch asset verified on the official website, 19 October 2026.
    if 'BTA-ALBORI' in projects:
        projects['BTA-ALBORI'].update(image='https://www.constructoracapital.com/web_datas/1790270330_peatonal-acceso-principal.webp',url='https://www.constructoracapital.com/lpage/266/albori/16',zone='Lagos de Torca',kind='VIS')
    if 'MDE-ROSACAMPESTRE' in projects:
        projects['MDE-ROSACAMPESTRE'].update(image='https://www.constructoracapital.com/web_datas/1726690216_2024-landing-rosacampestre-mesa-de-trabajo-1.webp',url='https://www.constructoracapital.com/lcom/13/rosa-campestre',zone='Rionegro')
    if 'MDE-BOSQUEDEPIEDRA' in projects:
        projects['MDE-BOSQUEDEPIEDRA'].update(image='https://bosquepiedra.co/wp-content/uploads/2025/07/Artboard-2-copy-6.jpg',url='https://bosquepiedra.co/',zone='Robledo · Pajarito')
    # Longest match prevents Torre Estación C. C. and Urbania variants being mixed.
    candidates=[]
    for pid,p in projects.items():
        candidates.append((key(p['name']),pid))
    for a,b in aliases.items():
        for pid in projects:
            if pid.endswith('-'+b):candidates.append((a,pid))
    candidates.sort(key=lambda x:len(x[0]),reverse=True)
    def match(name,c):
        n=norm(name)
        if any(x in n for x in ['MARCA','SEGUIDORES']):return None
        for k,pid in candidates:
            if pid.startswith(c+'-') and k in n:return pid
        return None
    source_counts={}; excluded=collections.Counter(); unmatched=collections.Counter(); agg={}; campaigns={}
    def put(pid,ch,m,sp,imp,clk,conv,count,src):
        k=(pid,ch,m)
        if k not in agg:agg[k]=[0,0,0,0,0]
        for i,v in enumerate([sp,imp,clk,conv,count]):agg[k][i]+=v
        source_counts[src]=source_counts.get(src,0)+count
    # Reconstruct total campaign spend from outcome allocations. Repeated metrics
    # are counted once per distinct impression/click signature within campaign-month.
    mg=collections.defaultdict(list)
    for ri,r in enumerate(w['ReporteMeta'].iter_rows(min_row=2,values_only=True),2):
        if len(r)<15 or not isinstance(r[3],(int,float)) or not isinstance(r[2],(int,float)):continue
        if r[5]!='OUTCOME_LEADS':excluded['Meta: objetivo distinto de leads']+=1;continue
        if num(r[14])<=0 and num(r[12])<=0:continue
        c=city(r[1]); m=f'{int(r[3]):04d}-{int(r[2]):02d}'; name=str(r[4]);pid=match(name,c)
        if not pid:unmatched['Meta | '+name]+=1;continue
        mg[(pid,m,name)].append(r)
    meta_signature_groups=0
    for (pid,m,name),rs in mg.items():
        signatures={(num(r[7]),num(r[10])) for r in rs}
        meta_signature_groups+=len(rs)-len(signatures)
        sp=sum(num(r[14]) for r in rs); leads=sum(num(r[12]) for r in rs if r[6]=='Leads')
        put(pid,'meta',m,sp,sum(i for i,c in signatures),sum(c for i,c in signatures),leads,len(rs),'ReporteMeta')
        campaigns[(pid,'meta',name)]=m
    for sn in ['Google_Bogota','Google_Medellin']:
        for r in w[sn].iter_rows(min_row=2,values_only=True):
            if len(r)<15 or not isinstance(r[3],(int,float)) or not isinstance(r[2],(int,float)):continue
            if r[5] not in ['SEARCH','PERFORMANCE_MAX']:excluded['Google: video, display o demanda']+=1;continue
            if num(r[14])<=0 and num(r[12])<=0:continue
            c=city(r[1]);name=str(r[4]);pid=match(name,c)
            if not pid:unmatched['Google | '+name]+=1;continue
            m=f'{int(r[3]):04d}-{int(r[2]):02d}'
            put(pid,'google',m,num(r[14]),num(r[7]),num(r[10]),num(r[12]),1,sn);campaigns[(pid,'google',name)]=m
    # Select one TikTok sheet per city/month. Records inside an export are additive;
    # there is no daily key, so no fabricated daily attribution or deduplication.
    tiktok_sheets=['Tiktok Año Capital Bogotá','Tiktok Año Capital Medellín Cap','Diciembre Tiktokreport','Tiktokreport','Tiktok Reporte Marzo',' Tiktok Reporte Abril',' Tiktok Reporte Mayo',' Tiktok Reporte Junio','Tiktok Reporte Julio','Tiktok Reporte Agosto','Tiktok Reporte Septiembre','Tiktok Reporte Octubre']
    tg={}
    for sn in tiktok_sheets:
        rows=list(w[sn].values)
        if not rows:continue
        special=sn=='Tiktok Año Capital Medellín Cap'
        by=collections.defaultdict(list)
        for r in rows[1:]:
            if len(r)<12 or not isinstance(r[2],(float,int)) or not isinstance(r[3],(float,int)):continue
            m=f'{int(r[3]):04d}-{int(r[2]):02d}';c=city(r[1]);by[(c,m)].append(r)
        for cm,rs in by.items():tg[cm]=(sn,rs,special)
    tiktok_month_rows={}
    for (c,m),(sn,rs,special) in tg.items():
        for r in rs:
            if special:sp,imp,clk,conv=num(r[11]),num(r[5]),num(r[8]),num(r[10]);lead=any(x in norm(r[4]) for x in ['LEAD','CONVERSION','CLIENTESPOTENCIALES'])
            else:
                if len(r)<15:continue
                sp,imp,clk,conv=num(r[14]),num(r[7]),num(r[10]),num(r[12]);lead=r[5]=='LEAD_GENERATION'
            if not lead:excluded['TikTok: objetivo distinto de leads']+=1;continue
            if sp<=0 and conv<=0:continue
            name=str(r[4]);pid=match(name,c)
            if not pid:unmatched['TikTok | '+name]+=1;continue
            put(pid,'tiktok',m,sp,imp,clk,conv,1,sn);campaigns[(pid,'tiktok',name)]=m
        tiktok_month_rows[c+' '+m]=sn
    # Add only observed campaign names as fallback template evidence.
    for (pid,ch,name),month in sorted(campaigns.items(),key=lambda kv:kv[1],reverse=True):
        p=projects[pid]
        if ch not in p['templates'] and name.startswith(('FB_','GLG_','TT_')):
            p['templates'][ch]={'name':name,'source':'Reporte de plataforma','period':month}
    records=[]
    for (pid,ch,m),v in sorted(agg.items()):records.append([pid,ch,m,*[round(x,6) for x in v[:4]],int(v[4])])
    for pid,p in projects.items():
        own=[r for r in records if r[0]==pid];p['months']=len(set(r[2] for r in own));p['lastMonth']=max((r[2] for r in own),default=None)
        if not p['token'] or p['token'] in ['VIVOPARK_1','VIVOPARK_2']:p['token']='VIVOPARK'
    meta={'sourceFile':pathlib.Path(workbook).name,'analyzedAt':'2026-10-19','latestMonth':'2026-10','completeThrough':'2026-09','recordColumns':['project','channel','month','spend','impressions','clicks','results','sourceRows'],'sourceRows':sum(source_counts.values()),'sourceCounts':source_counts,'excluded':dict(excluded),'metaRepeatedMetricRows':meta_signature_groups,'unmatchedCampaigns':len(unmatched),'unmatchedRows':sum(unmatched.values()),'invalidNames':invalid_names,'tiktokSources':tiktok_month_rows,'notes':['Meta: gasto total de campañas OUTCOME_LEADS; resultados solo de ubicación Leads. Impresiones y clics únicos por firma de métricas en campaña/mes, sin IDs de campaña en el archivo.','Google: conversiones atribuidas de Search y Performance Max. Se conservan decimales. No equivalen a personas únicas ni necesariamente a leads calificados.','TikTok: un único reporte por ciudad/mes; se suman sus filas. La exportación no incluye día ni ID de registro: granularidad y duplicados internos no son auditables.','Se excluyen marca, seguidores, awareness, video/display de Google, campañas sin proyecto reconocible y hojas consolidadas que repiten las fuentes.','Octubre de 2026 está incompleto. El periodo base predeterminado termina en septiembre.','La lista de proyectos activos procede del flow de octubre; no implica disponibilidad comercial.']}
    data={'meta':meta,'projects':sorted(projects.values(),key=lambda p:(not p['active'],p['city'],p['name'])),'records':records,'templates':templates}
    attach_consolidated(w,data)
    pathlib.Path(output).write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
    print(json.dumps({'projects':len(projects),'active':sum(p['active'] for p in projects.values()),'monthlyRecords':len(data['records']),'sourceRows':data['meta']['sourceRows'],'cityTotals':{c:{k:s[k] for k in ['rows','planned','actual']} for c,s in data['meta']['consolidated'].items()}},ensure_ascii=False))

if __name__=='__main__':build(*sys.argv[1:4])
